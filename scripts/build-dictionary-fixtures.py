#!/usr/bin/env python3
"""Rebuild the synthetic Excel data-dictionary fixtures.

These workbooks follow the SEARCH catalogue layout (category, blank index
column, field, type, format, acceptable values, existing values, required,
null). They are not copies of a study dictionary.

Requires openpyxl:  python3 scripts/build-dictionary-fixtures.py
"""

from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "configs" / "samples"

CERVICAL_HEADERS = [
    "Data Fields Category",
    None,
    "Data Fields",
    "Description / Definition",
    "Data Type",
    "Format / Character length",
    "Value Range/ Acceptable values",
    "Existing values in dataset",
    "Units of Measure",
    "Default Value",
    "Required?",
    "Accepts null value?",
]

OVARIAN_HEADERS = [
    "Data Field Category",
    None,
    "Data Field ",
    "Description / Definition",
    "Data Type",
    "Format / Character length",
    "Value Range/ Acceptable values",
    "Existing values in dataset",
    "Units of Measure",
    "Default Value",
    "Required?",
    "Accepts null value?",
]


def write_row(ws, row, values):
    for col, value in enumerate(values, start=1):
        if value is not None:
            ws.cell(row, col, value)


def build_demo(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Dictionary"
    write_row(ws, 1, CERVICAL_HEADERS)
    rows = [
        ["Screening", 1, "Enrolment date ", "Date the participant entered the cohort", "Date/Time",
         "DATE, format: MM-DD-YYYY", "2020-2024", "31/01/2020 - 02/15/2024", "NA", "NA", "Yes", "No"],
        [None, 2, "Age at visit", "Age in whole years at the visit", "Numerical", "Integer",
         "18-90", "19-95", "years", "NA", "Yes", "No"],
        [None, 3, "Assay result ", "Qualitative assay call", "Categorical", "VARCHAR( )",
         "Neg\nPos", "Neg\nPos\nBlanks", "NA", "NA", "No", "Yes"],
        [None, 4, "Assay result", "Repeated field name", "Categorical", "VARCHAR( )",
         "Neg\nPos", "Neg", "NA", "NA", "No", "Yes"],
        [None, 5, "Consent", "Whether consent was recorded", "Categorical", "Boolean, Values: Yes/No",
         "Yes\nNo", "Yes\nNo", "NA", "NA", "Yes", "Yes"],
        ["Notes", 6, "Comment", "Free-text note", "Free text", "VARCHAR(200)",
         "Free text", "Narrative", "NA", "NA", "No", "Yes"],
    ]
    for offset, values in enumerate(rows, start=2):
        write_row(ws, offset, values)
    ws.merge_cells("A2:A6")

    side = wb.create_sheet("Sheet1")
    for index, value in enumerate(["Neg", "Pos", "Equivocal"], start=1):
        side.cell(index, 1, value)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_clean(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Dictionary"
    write_row(ws, 1, OVARIAN_HEADERS)
    rows = [
        ["Cohort", 1, "Subject age", "Age in whole years", "Numerical", "Integer",
         "0-120", "1-90", "years", "NA", "Yes", "No"],
        [None, 2, "Site code", "Recruiting site", "Categorical", "VARCHAR( )",
         "A\nB", "A\nB", "NA", "NA", "No", "Yes"],
        ["Follow up", 3, "Visit date", "Date of the visit", "Date/Time", "DATE, format: MM-DD-YYYY",
         "2020-2024", "01-15-2020\n06-01-2021", "NA", "NA", "Yes", "No"],
    ]
    for offset, values in enumerate(rows, start=2):
        write_row(ws, offset, values)
    ws.merge_cells("A2:A3")
    wb.save(path)


if __name__ == "__main__":
    build_demo(OUT / "search-dictionary-demo.xlsx")
    build_clean(OUT / "search-dictionary-clean.xlsx")
    print(f"Wrote fixtures in {OUT}")
